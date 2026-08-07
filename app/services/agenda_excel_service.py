from io import BytesIO

from openpyxl import (Workbook,load_workbook)
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.config import settings
from app.integrations.google_drive import (
    download_file,
    find_file_in_folder,
    update_excel_file,
    upload_excel,
)
from app.models.appointment import Appointment
from app.models.patient import Patient


AGENDA_FILENAME = "agenda.xlsx"

AGENDA_HEADERS = [
    "ID de cita",
    "Fecha",
    "Hora inicio",
    "Hora fin",
    "Paciente",
    "Teléfono",
    "Estado",
    "Expediente",
    "Google Calendar",
]


def create_empty_agenda() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Agenda"

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:I1"

    for column_index, header in enumerate(
        AGENDA_HEADERS,
        start=1,
    ):
        cell = sheet.cell(
            row=1,
            column=column_index,
            value=header,
        )

        cell.font = Font(
            bold=True,
            color="FFFFFF",
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

    column_widths = {
        "A": 40,
        "B": 15,
        "C": 14,
        "D": 14,
        "E": 30,
        "F": 20,
        "G": 16,
        "H": 18,
        "I": 18,
    }

    for column, width in column_widths.items():
        sheet.column_dimensions[column].width = width

    sheet.row_dimensions[1].height = 24

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output.getvalue()

def ensure_agenda_exists() -> dict[str, str]:
    folder_id = settings.google_drive_agenda_folder_id

    if not folder_id:
        raise RuntimeError(
            "Falta GOOGLE_DRIVE_AGENDA_FOLDER_ID."
        )

    existing_file = find_file_in_folder(
        filename=AGENDA_FILENAME,
        folder_id=folder_id,
    )

    if existing_file is not None:
        return {
            "status": "existing",
            "file_id": existing_file["id"],
            "name": existing_file["name"],
            "web_view_link": existing_file.get(
                "webViewLink",
                "",
            ),
        }

    agenda_bytes = create_empty_agenda()

    uploaded_file = upload_excel(
        excel_bytes=agenda_bytes,
        filename=AGENDA_FILENAME,
        folder_id=folder_id,
    )

    return {
        "status": "created",
        "file_id": uploaded_file["file_id"],
        "name": uploaded_file["name"],
        "web_view_link": uploaded_file.get(
            "web_view_link",
            "",
        ),
    }

def add_appointment_to_agenda(
    agenda_bytes: bytes,
    appointment: Appointment,
    patient: Patient,
    expediente_url: str | None,
) -> bytes:
    workbook = load_workbook(
        BytesIO(agenda_bytes)
    )

    sheet = workbook["Agenda"]

    appointment_id = str(appointment.id)

    # Evitar duplicados
    for row in range(2, sheet.max_row + 1):
        current_id = sheet.cell(
            row=row,
            column=1,
        ).value

        if current_id == appointment_id:
            return agenda_bytes

    next_row = sheet.max_row + 1

    sheet.cell(
        row=next_row,
        column=1,
        value=appointment_id,
    )

    sheet.cell(
        row=next_row,
        column=2,
        value=appointment.starts_at.date(),
    )

    sheet.cell(
        row=next_row,
        column=3,
        value=appointment.starts_at.time(),
    )

    sheet.cell(
        row=next_row,
        column=4,
        value=appointment.ends_at.time(),
    )

    sheet.cell(
        row=next_row,
        column=5,
        value=patient.full_name,
    )

    sheet.cell(
        row=next_row,
        column=6,
        value=patient.phone,
    )

    sheet.cell(
        row=next_row,
        column=7,
        value=appointment.status.value,
    )

    expediente_cell = sheet.cell(
        row=next_row,
        column=8,
        value="Abrir expediente" if expediente_url else "",
    )

    if expediente_url:
        expediente_cell.hyperlink = expediente_url
        expediente_cell.style = "Hyperlink"

    sheet.cell(
        row=next_row,
        column=9,
        value="Pendiente",
    )

    # Formato
    sheet.cell(
        row=next_row,
        column=2,
    ).number_format = "dd/mm/yyyy"

    sheet.cell(
        row=next_row,
        column=3,
    ).number_format = "hh:mm"

    sheet.cell(
        row=next_row,
        column=4,
    ).number_format = "hh:mm"

    output = BytesIO()

    workbook.save(output)
    output.seek(0)

    return output.getvalue()

def sync_new_appointment_to_agenda(
    appointment: Appointment,
    patient: Patient,
    expediente_url: str | None,
) -> dict[str, str]:
    folder_id = settings.google_drive_agenda_folder_id

    if not folder_id:
        raise RuntimeError(
            "Falta GOOGLE_DRIVE_AGENDA_FOLDER_ID."
        )

    existing_file = find_file_in_folder(
        filename=AGENDA_FILENAME,
        folder_id=folder_id,
    )

    if existing_file is None:
        ensure_agenda_exists()

        existing_file = find_file_in_folder(
            filename=AGENDA_FILENAME,
            folder_id=folder_id,
        )

        if existing_file is None:
            raise RuntimeError(
                "No fue posible crear agenda.xlsx."
            )

    agenda_bytes = download_file(
        existing_file["id"]
    )

    updated_bytes = add_appointment_to_agenda(
        agenda_bytes=agenda_bytes,
        appointment=appointment,
        patient=patient,
        expediente_url=expediente_url,
    )

    return update_excel_file(
        file_id=existing_file["id"],
        excel_bytes=updated_bytes,
    )