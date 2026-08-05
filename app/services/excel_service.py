from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from app.models.appointment import Appointment
from app.models.intake import IntakeSubmission
from app.models.patient import Patient


SECTION_TITLES = {
    "consultation_reason": "Motivo de consulta",
    "current_condition": "Historia del padecimiento actual",
    "personal_history": "Antecedentes personales",
    "physical_activity": "Actividad física y hábitos",
    "functional_limitations": "Limitaciones funcionales",
    "alarm_signs": "Signos de alarma",
    "patient_goals": "Objetivos del paciente",
}


def format_value(value: Any) -> str:
    if value is None:
        return ""

    if isinstance(value, bool):
        return "Sí" if value else "No"

    if isinstance(value, list):
        return ", ".join(str(item) for item in value)

    return str(value)


def create_patient_workbook(
    patient: Patient,
    submission: IntakeSubmission,
    appointment: Appointment | None,
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Valoración"

    sheet.merge_cells("A1:B1")
    sheet["A1"] = "FICHA DE VALORACIÓN FISIOTERAPÉUTICA"
    sheet["A1"].font = Font(bold=True, size=14)
    sheet["A1"].alignment = Alignment(horizontal="center")

    current_row = 3

    general_data = [
        ("Nombre completo", patient.full_name),
        ("Edad", patient.age),
        ("Sexo", patient.sex),
        ("Ocupación", patient.occupation),
        ("Teléfono", patient.phone),
        ("Fecha de valoración", submission.assessment_date),
    ]

    sheet.cell(current_row, 1, "Datos generales")
    sheet.cell(current_row, 1).font = Font(bold=True)
    current_row += 1

    for label, value in general_data:
        sheet.cell(current_row, 1, label)
        sheet.cell(current_row, 2, format_value(value))
        current_row += 1

    current_row += 1

    for section_key, section_answers in submission.answers.items():
        title = SECTION_TITLES.get(
            section_key,
            section_key.replace("_", " ").title(),
        )

        sheet.cell(current_row, 1, title)
        sheet.cell(current_row, 1).font = Font(bold=True)
        current_row += 1

        for question, answer in section_answers.items():
            sheet.cell(
                current_row,
                1,
                question.replace("_", " ").capitalize(),
            )
            sheet.cell(
                current_row,
                2,
                format_value(answer),
            )
            current_row += 1

        current_row += 1

    sheet.cell(current_row, 1, "Cita")
    sheet.cell(current_row, 1).font = Font(bold=True)
    current_row += 1

    if appointment is None:
        sheet.cell(current_row, 1, "Cita solicitada")
        sheet.cell(current_row, 2, "No")
    else:
        sheet.cell(current_row, 1, "Cita solicitada")
        sheet.cell(current_row, 2, "Sí")
        current_row += 1

        sheet.cell(current_row, 1, "Inicio")
        sheet.cell(
            current_row,
            2,
            appointment.starts_at.isoformat(),
        )
        current_row += 1

        sheet.cell(current_row, 1, "Fin")
        sheet.cell(
            current_row,
            2,
            appointment.ends_at.isoformat(),
        )

    sheet.column_dimensions["A"].width = 45
    sheet.column_dimensions["B"].width = 60

    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True,
            )

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return output.getvalue()

import re
import unicodedata


def create_safe_filename(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    without_accents = normalized.encode(
        "ascii",
        "ignore",
    ).decode("ascii")

    safe_value = re.sub(
        r"[^a-zA-Z0-9_-]+",
        "_",
        without_accents.strip(),
    )

    return safe_value.strip("_") or "paciente"